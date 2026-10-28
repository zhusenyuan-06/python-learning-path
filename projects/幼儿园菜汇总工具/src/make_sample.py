# -*- coding: utf-8 -*-
"""
make_sample.py —— 生成模拟的「下个月送货单」，用来测试工具在新月份上的表现。

会做这几件事：
  · 日期换成目标月份的工作日
  · 数量、单价在真实数据基础上小幅浮动（模拟市场波动和用量变化）
  · 故意掺入几个新品名，用来验证「待补充」提示功能
  · 表尾放上正确的合计行，用来验证对账功能

用法：
    python src/make_sample.py --from "input/5月表.xlsx" --month 2026-06
"""

from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from loader import load_source  # noqa: E402

HEADER = ["日期", "编号", "品名及规格", "预订数", "单位", "发货数量",
          "实收数量", "", "单价（元）", "金额（元）", "", ""]

# 故意掺进来的新品名 —— 编号 9001 开头，保证 mapping 里没有
# 格式：(编号, 品名, 单位, 典型数量, 单价, 抽取权重)
# 其中「花胶」「响螺片」不含任何关键词，会被归到「待补充」——专门用来演示提醒功能
NEW_ITEMS = [
    ("9001", "澳洲大龙虾", "斤", 2.0, 380.0, 1),
    ("9002", "云南小瓜", "斤", 12.0, 3.8, 3),
    ("9003", "黑虎虾", "斤", 6.0, 52.0, 2),
    ("9004", "猫山王榴莲", "斤", 15.0, 32.0, 1),
    ("9005", "花胶", "斤", 1.0, 260.0, 1),        # 无关键词 -> 待补充
    ("9006", "响螺片", "斤", 2.0, 88.0, 1),       # 无关键词 -> 待补充
    ("9007", "三黄鸡", "斤", 18.0, 13.5, 2),
]

# 每个新品名至少出现的次数（保证演示时一定看得到）
MIN_NEW_TIMES = 2


def business_days(year: int, month: int, keep: int = 20) -> list[pd.Timestamp]:
    """目标月份的送货日：工作日里再均匀剔除几天（模拟停送/假期），保留 keep 天。"""
    start = pd.Timestamp(year=year, month=month, day=1)
    end = start + pd.offsets.MonthEnd(0)
    days = [d for d in pd.date_range(start, end) if d.weekday() < 5]
    if len(days) > keep:
        drop_count = len(days) - keep
        step = len(days) / drop_count
        drop_idx = {int(i * step) % len(days) for i in range(drop_count)}
        days = [d for i, d in enumerate(days) if i not in drop_idx]
    return days


def build_rows(df: pd.DataFrame, days: list[pd.Timestamp], per_day: tuple, rng: random.Random,
               new_items: list, seed_name: str):
    """基于真实明细，生成一个新月份的明细行。"""
    # 每个品名的典型用量与单价
    stat = df.groupby("品名").agg(
        次数=("日期", "nunique"),
        均量=("实收数量", "mean"),
        单价=("单价", "mean"),
        单位=("单位", "first"),
        编号=("编号", "first"),
    ).reset_index()

    # 高频菜品更容易出现，权重按出现次数
    names = list(stat["品名"])
    weights = [min(s, 12) for s in stat["次数"]]
    info = {r["品名"]: r for _, r in stat.iterrows()}

    # 新品名也加进候选池，权重低一些
    pool = names + [n[1] for n in new_items]
    pool_weights = weights + [n[5] for n in new_items]
    for code, name, unit, qty, price, _w in new_items:
        info[name] = {"编号": code, "单位": unit, "均量": qty, "单价": price}

    rows = []
    used_new = set()
    for day in days:
        n = rng.randint(*per_day)
        picked = rng.choices(pool, weights=pool_weights, k=n)
        for name in dict.fromkeys(picked):        # 同一天不重复
            it = info[name]
            base_qty = float(it["均量"])
            base_price = float(it["单价"])

            qty = round(base_qty * rng.uniform(0.85, 1.15), 1)
            qty = max(qty, 0.1)
            price = round(base_price * rng.uniform(0.94, 1.06), 2)

            # 预订数跟实收略有出入，跟真实单据一样
            ordered = round(qty * rng.uniform(0.98, 1.05), 1)

            amount = round(qty * price, 2)
            rows.append([
                day.to_pydatetime(), str(it["编号"]), name,
                ordered, str(it["单位"]), qty, qty, str(it["单位"]),
                price, amount, None, None,
            ])

    # 保证每个新品名都出现够 MIN_NEW_TIMES 次，不然演示时看不到
    for code, name, unit, base_qty, base_price, _w in new_items:
        times = sum(1 for r in rows if r[2] == name)
        while times < MIN_NEW_TIMES:
            day = rng.choice(days)
            qty = round(base_qty * rng.uniform(0.8, 1.2), 1)
            price = round(base_price * rng.uniform(0.95, 1.08), 2)
            rows.append([
                day.to_pydatetime(), code, name,
                round(qty * 1.02, 1), unit, qty, qty, unit,
                price, round(qty * price, 2), None, None,
            ])
            times += 1
        if times:
            used_new.add(name)

    rows.sort(key=lambda r: (r[0], int(r[1]) if str(r[1]).isdigit() else 9999))
    return rows, used_new


def write_workbook(path: Path, month_label: str, sheets: dict):
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.append([f"云湖湾{month_label}电子送货单"])
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)

        ws.append(HEADER)
        for c in ws[2]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")

        for r in rows:
            ws.append(r)

        total = round(sum(r[9] for r in rows), 2)
        ws.append([None] * 12)
        ws.append([None] * 9 + [total] + [None, None])       # 小计行
        ws.append([None] * 12)
        ws.append(["合计："] + [None] * 8 + [total] + [None, None])

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["I"].width = 10
        ws.column_dimensions["J"].width = 12

    wb.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--from", dest="src", help="参考的真实月份表")
    ap.add_argument("--month", default="2026-06", help="目标月份，格式 YYYY-MM")
    ap.add_argument("--out", default=None, help="输出路径，默认放 input 目录")
    ap.add_argument("--seed", type=int, default=202606)
    args = ap.parse_args()

    src = Path(args.src) if args.src else None
    if src is None:
        candidates = [p for p in (ROOT / "input").glob("*.xls*") if not p.name.startswith("~$")]
        if not candidates:
            print("input 目录里没有可参考的表。")
            return 1
        src = sorted(candidates)[0]     # 用最早的（真实月份）当参考

    year, month = (int(x) for x in args.month.split("-"))
    rng = random.Random(args.seed)

    real, _ = load_source(src)
    if not real:
        print(f"参考表里没读到明细：{src}")
        return 1

    days = business_days(year, month)
    out_sheets = {}
    injected = set()

    for sheet in real:
        per_day = (18, 24) if len(sheet.df) > 200 else (3, 5)
        new_items = NEW_ITEMS if len(sheet.df) > 200 else []
        rows, used_new = build_rows(sheet.df, days, per_day, rng, new_items, sheet.name)
        out_sheets[sheet.name] = rows
        injected |= used_new

    out = Path(args.out) if args.out else ROOT / "input" / f"云湖湾{year}.{month}月电子汇总表.xlsx"
    write_workbook(out, f"{year}.{month}月", out_sheets)

    print("─" * 60)
    print(f"已生成模拟表：{out}")
    print(f"送货日 {len(days)} 天（{days[0]:%Y-%m-%d} ~ {days[-1]:%Y-%m-%d}）")
    for name, rows in out_sheets.items():
        print(f"  {name}：{len(rows)} 行，合计 {sum(r[9] for r in rows):,.2f}")
    print(f"  注入新品名 {len(injected)} 个：{'、'.join(sorted(injected))}")
    print("─" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
