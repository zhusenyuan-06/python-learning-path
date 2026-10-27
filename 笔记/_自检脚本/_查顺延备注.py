import openpyxl
PATH = r"C:/Users/15816/Desktop/AI应用工程师学习日程表_已按摸底调整.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["每日计划"]
for r in (29, 30, 31, 32):
    d = str(ws.cell(r, 2).value)[:10]
    print(f"r{r} {d} {ws.cell(r,3).value}  时长={ws.cell(r,8).value}")
    for c in (5, 6, 7, 9):
        hdr = ws.cell(1, c).value
        print(f"    [{hdr}] {ws.cell(r,c).value}")
    print()
