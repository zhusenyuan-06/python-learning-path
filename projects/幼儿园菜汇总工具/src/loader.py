# -*- coding: utf-8 -*-
"""
loader.py —— 读取原始送货单并清洗成干净的明细表。

三条铁律：
1. 不依赖固定列位置：按表头名字匹配列，列顺序变了也不崩。
2. 不依赖 sheet 名：表里只要有"品名"表头就认，换个名字也能读。
3. 只读源文件，绝不修改它。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# 标准字段名 -> 可能出现的表头写法（第一个是最常见的）
COLUMN_ALIASES = {
    "日期": ["日期", "送货日期", "送货日", "配送日期"],
    "编号": ["编号", "商品编号", "货号", "货编号"],
    "品名": ["品名及规格", "品名规格", "品名", "名称", "商品名称", "货品名称"],
    "预订数": ["预订数", "预定数", "预订数量", "预定数量"],
    "单位": ["单位"],
    "发货数量": ["发货数量", "发货数", "发货"],
    "实收数量": ["实收数量", "实收数", "实收"],
    "单价": ["单价（元）", "单价(元)", "单价"],
    "金额": ["金额（元）", "金额(元)", "金额", "金额合计"],
}

# 这些行不是明细，要剔掉（但它们的金额要记下来用于对账）
SUMMARY_MARKERS = ("合计", "小计", "总计", "总金额")

OUTPUT_COLUMNS = [
    "日期", "编号", "品名", "单位",
    "预订数", "发货数量", "实收数量",
    "单价", "金额", "重算金额",
]


def _norm(text) -> str:
    """表头标准化：去空格、全角括号转半角。"""
    if text is None:
        return ""
    s = str(text).strip().replace(" ", "").replace("\u3000", "")
    return s.replace("（", "(").replace("）", ")")


def _to_float(value):
    """把单元格值转成 float，转不了返回 None。"""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _to_date(value):
    """转成 pandas 的 Timestamp，转不了返回 NaT。"""
    if value is None or value == "":
        return pd.NaT
    try:
        return pd.to_datetime(value)
    except (ValueError, TypeError):
        return pd.NaT


def find_header_row(ws, max_scan: int = 20):
    """在前 max_scan 行里找表头行：含有"品名"字样的那一行。"""
    for row_idx in range(1, min(max_scan, ws.max_row) + 1):
        for cell in ws[row_idx]:
            if "品名" in _norm(cell.value):
                return row_idx
    return None


def map_columns(header_cells: list) -> dict:
    """把表头单元格映射成 {标准字段名: 列索引}。"""
    normalized = [_norm(c) for c in header_cells]
    result = {}
    for std_name, aliases in COLUMN_ALIASES.items():
        for idx, text in enumerate(normalized):
            if not text:
                continue
            for alias in aliases:
                if _norm(alias) == text:
                    result.setdefault(std_name, idx)
                    break
            if std_name in result:
                break
    return result


@dataclass
class SheetData:
    """一个工作表清洗后的结果。"""
    name: str                      # sheet 名（幼儿 / 教师）
    df: pd.DataFrame               # 明细
    header_row: int                # 表头在第几行
    raw_row_count: int = 0         # 扫过的行数
    empty_rows: int = 0            # 跳过的空行
    summary_rows: int = 0          # 剔除的合计/小计行
    inner_totals: list = field(default_factory=list)   # 表内自带的金额数字
    warnings: list = field(default_factory=list)

    @property
    def total_amount(self) -> float:
        return float(self.df["重算金额"].sum()) if len(self.df) else 0.0


def read_sheet(ws) -> SheetData | None:
    """读一个工作表。没有品名表头就返回 None（比如"打印"这种模板表）。"""
    header_row = find_header_row(ws)
    if header_row is None:
        return None

    header_cells = [c.value for c in ws[header_row]]
    colmap = map_columns(header_cells)

    missing = [c for c in ("品名", "单价", "金额") if c not in colmap]
    if missing:
        return None  # 关键列不全，不是明细表

    # 数量列：优先「实收数量」，有些供应商只给「发货数量」
    qty_field = "实收数量" if "实收数量" in colmap else "发货数量"
    if qty_field not in colmap:
        return None

    name_idx = colmap["品名"]
    amt_idx = colmap["金额"]
    qty_idx = colmap[qty_field]
    price_idx = colmap["单价"]

    records = []
    inner_totals = []
    empty_rows = 0
    summary_rows = 0
    warnings = []

    for r in range(header_row + 1, ws.max_row + 1):
        values = [c.value for c in ws[r]]

        def cell(std_name):
            idx = colmap.get(std_name)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        raw_name = cell("品名")
        amount = _to_float(cell("金额"))

        # 空行
        if raw_name is None or str(raw_name).strip() == "":
            if amount is not None:
                inner_totals.append((r, amount))   # 孤立的金额，多半是表内小计
            else:
                empty_rows += 1
            continue

        name = str(raw_name).strip()

        # 合计/小计行
        if any(marker in name for marker in SUMMARY_MARKERS):
            if amount is not None:
                inner_totals.append((r, amount))
            summary_rows += 1
            continue

        qty = _to_float(cell(qty_field))
        price = _to_float(cell("单价"))

        # 分类标题行：表尾那种「蔬菜类」「豆制品」，只有名字，数量和金额都是空的
        # （绿禾那种没编号的明细行数量金额都有值，不会误伤）
        if qty is None and (amount is None or abs(amount) < 1e-9):
            empty_rows += 1
            continue

        # 数量或单价缺一个：有金额就直接采信表内金额，别算成 0
        use_table_amount = False
        if qty is None or price is None:
            if amount is not None:
                use_table_amount = True
                warnings.append(f"第{r}行「{name}」缺数量或单价，直接采用表内金额 {amount:.2f}")
            qty = qty or 0.0
            price = price or 0.0

        recalc = round(qty * price, 2)
        if use_table_amount:
            recalc = amount
        elif amount is None:
            amount = recalc
            warnings.append(f"第{r}行「{name}」金额为空，已用 数量×单价 补算")
        elif abs(recalc - amount) > 0.02:
            warnings.append(
                f"第{r}行「{name}」金额对不上：表内 {amount:.2f}，数量×单价 {recalc:.2f}，已按 数量×单价 计"
            )
            amount = recalc

        records.append({
            "日期": _to_date(cell("日期")),
            "编号": str(cell("编号")).strip() if cell("编号") is not None else "",
            "品名": name,
            "单位": str(cell("单位")).strip() if cell("单位") is not None else "",
            "预订数": _to_float(cell("预订数")),
            "发货数量": _to_float(cell("发货数量")),
            "实收数量": qty,
            "单价": price,
            "金额": amount,
            "重算金额": recalc,
        })

    df = pd.DataFrame(records, columns=OUTPUT_COLUMNS)
    if len(df):
        df["日期"] = pd.to_datetime(df["日期"])
        # 送货单里日期大多是合并单元格：同一天只在第一行写日期，下面留空。
        # 往下填一次，否则"送货次数""按日统计"全都会算错。
        df["日期"] = df["日期"].ffill()

    return SheetData(
        name=ws.title,
        df=df,
        header_row=header_row,
        raw_row_count=max(0, ws.max_row - header_row),
        empty_rows=empty_rows,
        summary_rows=summary_rows,
        inner_totals=inner_totals,
        warnings=warnings,
    )


def supplier_name(sheet_name: str, aliases: dict | None = None) -> str:
    """从 sheet 名里取供应商名：去掉尾部的月份，再查一遍别名表。

    「绿禾9月」-> 绿禾     「幼儿 9月」-> 幼儿 ->（别名）-> 锦源
    """
    import re
    s = str(sheet_name).strip()
    raw = s
    s = re.sub(r"[\s_\-]*\d{1,2}\s*月\s*$", "", s)        # 去掉尾部 9月
    s = re.sub(r"[\s_\-]*\d{4}[-./年]\d{1,2}\s*$", "", s)  # 去掉尾部 2026-09
    s = s.strip()
    if aliases:
        # 先用完整的 sheet 名查（这样可以只纠正某个月的写法，不影响历史表）
        s = str(aliases.get(raw, aliases.get(s, s))).strip()
    return s or str(sheet_name).strip()


def load_source(path: str | Path, min_rows: int = 3, supplier_aliases: dict | None = None):
    """读取一个工作簿里的所有明细表。

    返回 (明细表列表, 被跳过的表及其原因)。
    min_rows：明细少于这个数的表当成模板表跳过（比如只有一两行示例的"打印"表）。
    supplier_aliases：sheet 名 -> 供应商名的纠错表。
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets, skipped = [], []
        for ws in wb.worksheets:
            data = read_sheet(ws)
            if data is None:
                skipped.append((ws.title, "没有『品名』表头或关键列不全"))
                continue
            if len(data.df) < min_rows:
                skipped.append((ws.title, f"只有 {len(data.df)} 行明细，当成模板表跳过"))
                continue
            data.name = supplier_name(ws.title, supplier_aliases)
            sheets.append(data)
        return sheets, skipped
    finally:
        wb.close()
