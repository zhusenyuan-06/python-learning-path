# -*- coding: utf-8 -*-
"""敏感信息自检（只读，不修改）

这个脚本不依赖任何具体的企业名/园所名 —— 它按【通用模式】扫，
所以它自己也不会成为泄密点。扫的是这四类：

  1. 手机号         1[3-9] 开头 11 位
  2. 固定电话       0755-1234567 这种带区号的
  3. 企业全称       xx有限公司 / xx配送服务有限公司
  4. 详细地址       同时出现「省/市/区/街道/路/号/物流园」这类地名片段

白名单里的都是本仓库虚构的替换值，扫到不算问题。

用法：
    python _检查敏感数据.py
"""
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "src"))
from paths import get_root  # noqa: E402

ROOT = get_root()

# 本仓库使用的虚构值，扫到不算命中
白名单 = ["云湖湾", "锦源", "绿禾", "示例市", "0755-6688000", "13900001234"]

规则 = [
    ("手机号", re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")),
    ("固定电话", re.compile(r"(?<!\d)\d{3,4}-\d{7,8}(?!\d)")),
    ("企业全称", re.compile(r"[一-龥]{2,12}(?:有限责任公司|股份有限公司|有限公司|配送服务有限公司|配送中心|商贸行)")),
    ("详细地址", re.compile(r"[一-龥]{2,10}(?:省|市|区|县)[一-龥]{0,10}(?:街道|镇|路|道|号|物流园|工业园)[^\s，。；]{0,20}")),
]

TEXT_EXT = (".py", ".md", ".yaml", ".yml", ".csv", ".txt", ".bat", ".spec")
EXCEL_EXT = (".xlsx", ".xls")
# 自检脚本本身含有规则示例文本，扫描时排除自己
SELF = os.path.basename(os.path.abspath(__file__))


def 扫文本(t):
    out = []
    for 名, rx in 规则:
        for m in rx.findall(t):
            if any(w in m for w in 白名单):
                continue
            out.append((名, m))
    return out


def 收集文本(p):
    try:
        with open(p, encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return None


def main():
    print("=" * 70)
    print("敏感信息自检（通用模式，不依赖具体企业名）")
    print("=" * 70)

    命中 = []
    检查文件 = 0
    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if d not in ("__pycache__", ".git", "_构建产物")]
        for f in filenames:
            full = os.path.join(dirpath, f)
            rel = os.path.relpath(full, ROOT)
            if f == SELF:
                continue
            if f.endswith(TEXT_EXT):
                t = 收集文本(full)
                if t is None:
                    continue
                检查文件 += 1
                for 名, 片段 in 扫文本(t):
                    命中.append((rel, 名, 片段))
            elif f.endswith(EXCEL_EXT):
                try:
                    wb = openpyxl.load_workbook(full, data_only=True)
                except Exception:
                    continue
                检查文件 += 1
                for sn in wb.sheetnames:
                    for row in wb[sn].iter_rows():
                        for c in row:
                            if isinstance(c.value, str):
                                for 名, 片段 in 扫文本(c.value):
                                    命中.append(("%s[%s]%s" % (rel, sn, c.coordinate), 名, 片段))

    print("已扫描 %d 个文件" % 检查文件)
    print()
    if 命中:
        print("[!] 发现 %d 处疑似敏感内容：" % len(命中))
        for h in 命中[:25]:
            print("     ", h)
    else:
        print("[OK] 未发现手机号 / 固话 / 企业全称 / 详细地址")
        print("     （白名单内的虚构值已排除）")

    print()
    print("=" * 70)
    print("附：脱敏后的替换对照（虚构值）")
    print("=" * 70)
    for 名, 值 in [("园所名", "云湖湾（虚构）"), ("供应商 A", "锦源（虚构）"),
                    ("供应商 B", "绿禾（虚构）"), ("固话", "0755-6688000（虚构）"),
                    ("手机", "13900001234（虚构）"), ("地址", "示例市城东区…（虚构）")]:
        print("  %-10s %s" % (名, 值))
    return 1 if 命中 else 0


if __name__ == "__main__":
    raise SystemExit(main())
