import openpyxl

PATH = r"C:\Users\15816\Desktop\AI应用工程师学习日程表_已按摸底调整.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["每日计划"]

targets = {"2026-10-27", "2026-10-28", "2026-10-30", "2026-11-08"}

print("=" * 100)
print("【今日 10-27 及相关日期在日程表里的原定安排】")
print("=" * 100)
for r in range(2, ws.max_row + 1):
    d = ws.cell(r, 2).value
    ds = str(d)[:10] if d is not None else ""
    if ds in targets:
        print(f"\n--- r{r}  日期={ds}  {ws.cell(r,3).value}  {ws.cell(r,4).value}  时长={ws.cell(r,8).value}h  完成={ws.cell(r,9).value!r}")
        for col, label in [(5, "上午"), (6, "下午"), (7, "晚上")]:
            v = ws.cell(r, col).value
            print(f"    {label}: {v}")

print()
print("=" * 100)
print("【全部含 *args / kwargs / 导入 / 模块 关键词的行】")
print("=" * 100)
kw = ["args", "kwargs", "导入", "模块"]
for r in range(2, ws.max_row + 1):
    row_text = " | ".join(str(ws.cell(r, c).value) for c in range(1, 10) if ws.cell(r, c).value is not None)
    if any(k in row_text for k in kw):
        print(f"r{r}  {ws.cell(r,2).value} {ws.cell(r,3).value} | {row_text[:150]}")
