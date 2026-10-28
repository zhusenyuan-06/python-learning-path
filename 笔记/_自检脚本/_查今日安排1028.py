import openpyxl
PATH = r"C:/Users/15816/Desktop/AI应用工程师学习日程表_已按摸底调整.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["每日计划"]
print("=" * 96)
for r in range(2, ws.max_row + 1):
    d = str(ws.cell(r, 2).value)[:10]
    if d in ("2026-10-27", "2026-10-28", "2026-10-29"):
        print(f"\n--- r{r}  {d}  {ws.cell(r,3).value}  {ws.cell(r,4).value}  时长={ws.cell(r,8).value}h")
        print(f"    完成列: {ws.cell(r,9).value!r}")
        for c, lab in [(5,"上午"),(6,"下午"),(7,"晚上")]:
            print(f"    [{lab}] {ws.cell(r,c).value}")
print("\n" + "=" * 96)
