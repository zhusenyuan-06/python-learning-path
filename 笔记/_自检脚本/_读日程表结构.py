import pandas as pd
import openpyxl

PATH = r"C:\Users\15816\Desktop\AI应用工程师学习日程表_已按摸底调整.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

print("sheets:", wb.sheetnames)
print()

for name in wb.sheetnames:
    ws = wb[name]
    print("=" * 70)
    print(f"[sheet] {name}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    print("=" * 70)
    for r in range(1, min(ws.max_row, 6) + 1):
        vals = [str(ws.cell(r, c).value)[:28] if ws.cell(r, c).value is not None else "" for c in range(1, min(ws.max_column, 10) + 1)]
        print(f"  r{r}: {vals}")
    print()
