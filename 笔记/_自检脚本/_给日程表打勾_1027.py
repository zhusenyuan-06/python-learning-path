# -*- coding: utf-8 -*-
"""给日程表 2026-10-27 那一行打勾（幂等，可重复运行）"""
import datetime
import os
import shutil
import openpyxl

PATH = r"C:\Users\15816\Desktop\AI应用工程师学习日程表_已按摸底调整.xlsx"
备份目录 = r"D:\AI应用工程师\09-求职资产"

def main():
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    备份 = os.path.join(备份目录, f"日程表.打勾备份_{stamp}.xlsx")
    os.makedirs(备份目录, exist_ok=True)
    shutil.copy2(PATH, 备份)
    print("✅ 已备份:", 备份)

    wb = openpyxl.load_workbook(PATH)
    ws = wb["每日计划"]
    目标日期 = "2026-10-27"
    改了几格 = 0
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 2).value
        ds = str(d)[:10] if d is not None else ""
        if ds == 目标日期:
            旧 = ws.cell(r, 9).value
            ws.cell(r, 9).value = "✅ 已完成"
            改了几格 += 1
            print(f"  r{r} {ds} 完成列: {旧!r} → '✅ 已完成'")
            print(f"     时长: {ws.cell(r,8).value}h")
    if 改了几格 == 0:
        print("⚠️ 没找到对应行，未修改")
        return
    wb.save(PATH)
    print(f"✅ 已保存，共改 {改了几格} 格")

    # 回读校验
    wb2 = openpyxl.load_workbook(PATH, data_only=True)
    ws2 = wb2["每日计划"]
    for r in range(2, ws2.max_row + 1):
        d = ws2.cell(r, 2).value
        if str(d)[:10] == 目标日期:
            print("✅ 回读校验 完成列 =", repr(ws2.cell(r, 9).value))
    print("✅ sheet 列表:", wb2.sheetnames)

if __name__ == "__main__":
    main()
